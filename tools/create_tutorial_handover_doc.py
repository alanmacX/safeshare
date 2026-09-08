from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
from datetime import datetime, timezone
import shutil
import tempfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "entry/src/main/resources/rawfile/project_handover.docx"


def set_run_font(run, size=11, bold=False, color="222222"):
    run.font.name = "Heiti SC"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Heiti SC")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{margin}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def add_label_value(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    label_run = p.add_run(label)
    set_run_font(label_run, 10, True, "555555")
    value_run = p.add_run(value)
    set_run_font(value_run, 10, False, "222222")


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Heiti SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title_style_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_style_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_style_ppr.remove(title_border)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(16)
    run = title.add_run("Project Delivery Confirmation")
    set_run_font(run, 22, True, "000000")

    add_label_value(doc, "Project    ", "Harbor Center Wayfinding Upgrade")
    add_label_value(doc, "Client    ", "Galaxy Commercial Management Co., Ltd.")
    add_label_value(doc, "Supplier    ", "Yuanlan Digital Technology Co., Ltd.")
    add_label_value(doc, "Delivery date    ", "28 August 2026")

    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(7)
    set_run_font(heading.add_run("Deliverables"), 13, True, "000000")

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Cm(1.3), Cm(9.4), Cm(3.3)]
    headers = ["No.", "Deliverable", "Status"]
    for index, cell in enumerate(table.rows[0].cells):
        cell.width = widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "244E46")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(headers[index]), 10, True, "FFFFFF")

    rows = [
        ("1", "Wayfinding source files and font inventory", "Delivered"),
        ("2", "Production dimensions and material specification", "Delivered"),
        ("3", "Installation map and acceptance photo archive", "Delivered"),
    ]
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].width = widths[index]
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[index], top=130, bottom=130)
            if row_index % 2 == 1:
                set_cell_shading(cells[index], "F3F7F6")
            p = cells[index].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index != 1 else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(value), 10, False, "222222")

    for heading_text, body in [
        ("Delivery confirmation", "Both parties confirm that the deliverables listed above are complete and usable. New locations and construction changes require separate approval."),
        ("Information handling", "Source files are provided only for production and maintenance of this project. Personal names, internal paths and revision metadata must be removed before external distribution."),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(13)
        p.paragraph_format.space_after = Pt(5)
        set_run_font(p.add_run(heading_text), 13, True, "000000")
        body_p = doc.add_paragraph()
        body_p.paragraph_format.line_spacing = 1.35
        body_p.paragraph_format.space_after = Pt(4)
        set_run_font(body_p.add_run(body), 10.5, False, "333333")

    sign = doc.add_table(rows=2, cols=2)
    sign.autofit = False
    sign.columns[0].width = Cm(7)
    sign.columns[1].width = Cm(7)
    labels = ["Client representative: _______________", "Supplier representative: _______________",
              "Date: __________________", "Date: __________________"]
    for index, cell in enumerate([c for row in sign.rows for c in row.cells]):
        set_cell_margins(cell, top=180, bottom=180)
        set_run_font(cell.paragraphs[0].add_run(labels[index]), 10, False, "333333")

    props = doc.core_properties
    props.title = "Harbor Center Wayfinding Upgrade Project Delivery Confirmation"
    props.subject = "Project delivery and acceptance"
    props.author = "Jia Ning Chen"
    props.last_modified_by = "Mingyuan Zhou"
    props.keywords = "Harbor Center;wayfinding;delivery;acceptance"
    props.comments = "Client approval copy containing contact and internal revision metadata"
    props.created = datetime(2026, 8, 20, 9, 15, tzinfo=timezone.utc)
    props.modified = datetime(2026, 8, 28, 16, 40, tzinfo=timezone.utc)
    doc.save(OUTPUT)

    with tempfile.TemporaryDirectory() as temp_dir:
        temp = Path(temp_dir)
        with ZipFile(OUTPUT, "r") as zin:
            zin.extractall(temp)
        app_xml = temp / "docProps/app.xml"
        text = app_xml.read_text(encoding="utf-8")
        text = text.replace("</Properties>",
            "<Company>Yuanlan Digital Technology Co., Ltd.</Company><Manager>Zhiyuan Lin</Manager></Properties>")
        app_xml.write_text(text, encoding="utf-8")
        rebuilt = OUTPUT.with_suffix(".rebuilt.docx")
        with ZipFile(rebuilt, "w", ZIP_DEFLATED) as zout:
            for path in sorted(temp.rglob("*")):
                if path.is_file():
                    zout.write(path, path.relative_to(temp).as_posix())
        shutil.move(rebuilt, OUTPUT)


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    build_docx()
